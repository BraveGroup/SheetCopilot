"""
Open-loop planning probe for SheetCopilot.

The full SheetCopilot agent is *closed-loop*: it plans a step, executes it in
Excel via ``win32com``, observes the new sheet state, and repeats.  That requires
Windows + Excel and therefore cannot run on Linux.

This probe lets you **test any OpenAI-compatible model on N tasks without Excel**.
For each task it builds the real SheetCopilot planning prompt (the system prompt +
action-API document + the few-shot example from ``config/prompt.yaml``) plus a
sheet-state description read from the source workbook with openpyxl, asks the
model for a complete step-by-step plan, and saves a full trajectory (the query,
the model response + reasoning, token usage, latency, and the parsed/validated
action APIs).  It measures *planning* quality (does the model produce valid
SheetCopilot action plans?), not end-to-end Pass@1 (which needs Excel execution).

It exercises the modernized OpenAI client (``utils/ChatGPT.py``) and trajectory
logger (``utils/trajectory.py``) against a live endpoint.

Example (point at any OpenAI-compatible server)::

    python run_planning_probe.py \
        --base-url https://host/path/v1 --api-key KEY --model MODEL \
        --num 10 --task-file ../dataset/dataset_20Samples.xlsx

Credentials may also come from the OPENAI_BASE_URL / OPENAI_API_KEY / OPENAI_MODEL
environment variables.
"""

import argparse
import asyncio
import importlib.util
import json
import os
import re
import sys
import time

import openpyxl
import pandas as pd
import yaml

HERE = os.path.dirname(os.path.abspath(__file__))


def _load(name, relpath):
    """Import a single module file directly (bypassing utils/__init__, which
    imports the Windows-only win32com-based comparison code)."""
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, relpath))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod

_cg = _load("sc_chatgpt", "utils/ChatGPT.py")
_tj = _load("sc_traj", "utils/trajectory.py")
_cp = _load("sc_cp", "utils/construct_prompt.py")

ChatGPT = _cg.ChatGPT
TrajectoryLogger = _tj.TrajectoryLogger

# SheetCopilot @API(...)@ form, mirroring agent/Agent/agent.py.
RE_FULL_CALL = re.compile(r'(?<=@)([A-Z].*?\))(?=@|\n|$)')
# Qwen3 / Hermes tool-call form some models emit:
#   <tool_call><function=Write><parameter=range>A1</parameter>...</function></tool_call>
RE_QWEN_FUNC = re.compile(r'<function\s*=\s*([A-Za-z0-9_]+)\s*>(.*?)</function>', re.S)
RE_QWEN_PARAM = re.compile(r'<parameter\s*=\s*([A-Za-z0-9_]+)\s*>\s*(.*?)\s*</parameter>', re.S)
RE_CALL_NAME = re.compile(r'^\s*([A-Za-z_][A-Za-z0-9_]*)\s*\(')


def parse_qwen_tool_calls(text):
    """Convert Qwen3-style <tool_call> blocks into ``Name(arg="val", ...)`` strings."""
    calls = []
    for name, body in RE_QWEN_FUNC.findall(text or ""):
        params = RE_QWEN_PARAM.findall(body)
        argstr = ", ".join('{}="{}"'.format(k, v.strip().replace('\n', ' ')) for k, v in params)
        calls.append("{}({})".format(name, argstr))
    return calls


def describe_workbook(path):
    """Build a SheetCopilot-style sheet-state string from a workbook."""
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for si, ws in enumerate(wb.worksheets):
        max_col = ws.max_column or 0
        # contiguous header cells from A1
        headers = []
        for c in range(1, max_col + 1):
            v = ws.cell(row=1, column=c).value
            if v is None:
                break
            headers.append((openpyxl.utils.get_column_letter(c), str(v)))
        ncols = len(headers) if headers else max_col
        nrows = ws.max_row or 0
        active = ' (active)' if si == 0 else ''
        if headers:
            hdr = ', '.join(f'{col}: "{name}"' for col, name in headers)
            parts.append(f'Sheet "{ws.title}"{active} has {ncols} columns (Headers are {hdr}) '
                         f'and {nrows} rows (the row 1 is the header row).')
        else:
            parts.append(f'Sheet "{ws.title}"{active} has {ncols} columns and {nrows} rows.')
    return 'Sheet state: ' + ' '.join(parts)


def build_messages(prompt_template, api_usage, context, instruction, sheet_state):
    """Compose the planning messages: system (with API doc) + few-shot + task.

    The final user turn asks for the *complete* plan in one response (the probe
    is open-loop, so there is no execution feedback between steps)."""
    msgs = [dict(m) for m in prompt_template]
    msgs[0] = dict(msgs[0])
    msgs[0]['content'] = msgs[0]['content'].format(API_Doc=api_usage)
    # Replace the final templated user turn with a complete-plan request.
    final = (
        f"{context}\n"
        f"Instruction: {instruction}\n"
        f"{sheet_state}\n"
        "Please provide the COMPLETE step-by-step solution in a SINGLE message. Output each step as:\n"
        "Step X. <brief reason>\n"
        "Action API: @SomeAPI(...)@\n"
        "Rules: use only the action APIs from the document above; wrap every API call in @...@ "
        "(do NOT use <tool_call>, <function=...> or any XML/JSON tool-call format); include the "
        "sheet name in every range; list ALL steps now without waiting for any tool result; and "
        "write \"Done!\" on the last line. Keep it concise without extra explanations."
    )
    msgs[-1] = {"role": "user", "content": final}
    return msgs


def parse_plan(content, api_list):
    """Extract action calls from a response in *either* the SheetCopilot
    ``@API(...)@`` form or the Qwen3 ``<tool_call>`` form (this model emits both),
    de-duplicate them, and validate the API names against the action space."""
    content = content or ""
    api_lower = {a.lower(): a for a in api_list}

    raw_calls = RE_FULL_CALL.findall(content) + parse_qwen_tool_calls(content)

    # De-duplicate (the model often repeats the same action in both forms).
    calls, seen = [], set()
    for c in raw_calls:
        key = re.sub(r"\s+", "", c)
        if key and key not in seen:
            seen.add(key)
            calls.append(c)

    valid, invalid = [], []
    for c in calls:
        m = RE_CALL_NAME.match(c)
        if not m:
            continue
        name = m.group(1)
        if name.lower() in api_lower:
            valid.append(api_lower[name.lower()])
        else:
            invalid.append(name)

    used_qwen = "<tool_call>" in content or "<function=" in content
    return {
        "n_steps": len(re.findall(r'(?im)^\s*step\s+\d+', content)) or len(calls),
        "n_calls": len(calls),
        "calls": calls,
        "valid_apis": valid,
        "invalid_apis": invalid,
        "done": "Done" in content,
        "used_qwen_toolcall": used_qwen,
    }


async def run_task(sem, args, prompt_template, api_usage, api_list, idx, row, out_dir):
    sheet_name = row["Sheet Name"]
    task_name = f"{idx + 1}_{sheet_name}"
    source = os.path.join(args.source_dir, sheet_name + ".xlsx")
    context = str(row.get("Context", "") or "")
    instruction = str(row.get("Instructions", "") or "")

    try:
        sheet_state = describe_workbook(source)
    except Exception as e:
        sheet_state = "Sheet state: (could not read source workbook: %s)" % e

    messages = build_messages(prompt_template, api_usage, context, instruction, sheet_state)

    cfg = {
        "base_url": args.base_url,
        "api_key": args.api_key,
        "model_name": args.model,
        "max_tokens": args.max_tokens,
        "temperature": args.temperature,
        "timeout": args.timeout,
        "max_total_tokens": 10 ** 9,
        "max_retries": 3,
        "sleep_time": 5,
        "stream": args.stream,
    }

    async with sem:
        logger = TrajectoryLogger(meta={
            "task": task_name, "model": args.model, "sheet_name": sheet_name,
            "instruction": instruction, "context": context, "source_file": source,
            "mode": "open_loop_planning_probe",
        })
        # context excludes the final user turn; pass it so the wrapper appends it.
        bot = ChatGPT(cfg, context=messages[:-1], logger=logger)
        bot.stage = "planning"
        t0 = time.time()
        error = None
        try:
            content = await bot(messages[-1]["content"])
        except Exception as e:
            content, error = None, str(e)
        latency = time.time() - t0

    parsed = parse_plan(content, api_list) if content else {
        "n_steps": 0, "n_calls": 0, "calls": [], "valid_apis": [], "invalid_apis": [], "done": False}
    logger.annotate_last(parsed_actions=parsed["calls"],
                         valid_apis=parsed["valid_apis"],
                         invalid_apis=parsed["invalid_apis"])

    task_dir = os.path.join(out_dir, task_name)
    traj_path = os.path.join(task_dir, f"{task_name}_trajectory.json")
    logger.save(traj_path, success=(error is None and parsed["n_calls"] > 0), error=error)

    usage = logger.calls[-1]["usage"] if logger.calls else None
    result = {
        "task": task_name,
        "instruction": instruction[:80],
        "latency_s": round(latency, 1),
        "total_tokens": (usage or {}).get("total_tokens"),
        "n_steps": parsed["n_steps"],
        "n_calls": parsed["n_calls"],
        "n_valid": len(parsed["valid_apis"]),
        "n_invalid": len(parsed["invalid_apis"]),
        "invalid_apis": sorted(set(parsed["invalid_apis"])),
        "done": parsed["done"],
        "error": error,
        "trajectory": os.path.relpath(traj_path, HERE),
    }
    flag = "ok " if (error is None and parsed["n_calls"] > 0 and not parsed["invalid_apis"]) else "!! "
    print(f"  {flag}{task_name:<28} steps={result['n_steps']:<2} apis={result['n_calls']:<2} "
          f"valid={result['n_valid']:<2} invalid={result['n_invalid']:<2} "
          f"{result['latency_s']:>5.1f}s tok={result['total_tokens']}"
          + (f"  ERROR: {error[:60]}" if error else ""))
    return result


async def main_async(args):
    with open(os.path.join(HERE, "config/prompt.yaml")) as f:
        prompt = yaml.load(f, Loader=yaml.Loader)["task planning"]
    with open(os.path.join(HERE, "config/API_document.yaml")) as f:
        api_doc = yaml.load(f, Loader=yaml.FullLoader)
    api_list, api_usage, _ = _cp.get_api_doc("gpt-chat-prompt", api_doc)

    df = pd.read_excel(args.task_file, header=0).head(args.num)
    out_dir = args.out or os.path.join(HERE, f"planning_probe_{args.model.replace('/', '_')}")
    os.makedirs(out_dir, exist_ok=True)

    print(f"Planning probe: model={args.model}  base_url={args.base_url}")
    print(f"  tasks={len(df)}  task_file={args.task_file}  out={out_dir}  stream={args.stream}\n")

    sem = asyncio.Semaphore(args.workers)
    jobs = [run_task(sem, args, prompt, api_usage, api_list, i, row, out_dir)
            for i, row in df.iterrows()]
    results = await asyncio.gather(*jobs)

    n = len(results)
    ok = sum(1 for r in results if r["error"] is None and r["n_calls"] > 0)
    clean = sum(1 for r in results if r["error"] is None and r["n_calls"] > 0 and r["n_invalid"] == 0)
    tot_tok = sum((r["total_tokens"] or 0) for r in results)
    avg_lat = sum(r["latency_s"] for r in results) / n if n else 0

    summary = {
        "model": args.model, "base_url": args.base_url, "task_file": args.task_file,
        "num_tasks": n, "produced_plan": ok, "plan_only_valid_apis": clean,
        "total_tokens": tot_tok, "avg_latency_s": round(avg_lat, 1),
        "results": results,
    }
    with open(os.path.join(out_dir, "summary.json"), "w") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    print("\n================ Planning probe summary ================")
    print(f"  model                       : {args.model}")
    print(f"  tasks                       : {n}")
    print(f"  produced a plan (>=1 API)   : {ok}/{n}")
    print(f"  plan with only valid APIs   : {clean}/{n}")
    print(f"  avg latency                 : {avg_lat:.1f}s")
    print(f"  total tokens                : {tot_tok}")
    print(f"  trajectories + summary.json : {out_dir}")
    print("  NOTE: open-loop planning quality only; end-to-end Pass@1 needs Excel execution.")
    return 0


def main():
    p = argparse.ArgumentParser(description="Open-loop SheetCopilot planning probe.")
    p.add_argument("--base-url", default=os.environ.get("OPENAI_BASE_URL"))
    p.add_argument("--api-key", default=os.environ.get("OPENAI_API_KEY", "EMPTY"))
    p.add_argument("--model", default=os.environ.get("OPENAI_MODEL", "gpt-4o-mini"))
    p.add_argument("--num", type=int, default=10)
    p.add_argument("--task-file", default=os.path.join(HERE, "..", "dataset", "dataset_20Samples.xlsx"))
    p.add_argument("--source-dir", default=os.path.join(HERE, "..", "dataset", "task_sheets"))
    p.add_argument("--out", default=None)
    p.add_argument("--workers", type=int, default=5)
    p.add_argument("--max-tokens", type=int, default=8000)
    p.add_argument("--temperature", type=float, default=0.0)
    p.add_argument("--timeout", type=int, default=300)
    p.add_argument("--stream", action="store_true")
    args = p.parse_args()
    if not args.base_url:
        p.error("--base-url is required (or set OPENAI_BASE_URL)")
    raise SystemExit(asyncio.run(main_async(args)))


if __name__ == "__main__":
    main()
