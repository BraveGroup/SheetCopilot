"""
Claude Code as an alternative SheetCopilot agent (Linux, Excel-free).

Instead of the built-in SheetCopilot state machine (which drives Excel via
``win32com`` on Windows), this harness hands each task to the **Claude Code CLI**
and lets it manipulate the workbook with ordinary tools (Python/openpyxl/pandas,
or headless LibreOffice).  The results are written in the exact folder/log layout
that ``evaluation_ubuntu.py`` expects, so the two agents are scored identically,
and a ``*_trajectory.json`` is emitted so the runs can be inspected with
``visualize_trajectories.py``.

For every task it:
  1. creates ``<save_path>/<order>_<Sheet>/`` and copies the source workbook to
     ``..._source.xlsx`` and to the working result ``..._1.xlsx``;
  2. runs ``claude -p <prompt> --output-format json`` in that folder, pointed at
     the source/result files, to edit the result in place;
  3. writes an evaluator-compatible ``..._log.yaml`` and a trajectory JSON.

Then evaluate with::

    python evaluation_ubuntu.py -c <config-with-this-save_path>

IMPORTANT — running Claude Code autonomously executes shell commands without
interactive approval, so:
  * it needs ``--permission-mode bypassPermissions`` (or an allow-list), which
    Claude Code refuses to run as **root** -- run this harness as a NON-root user;
  * set the model endpoint via the standard Claude Code env vars, e.g.::
        export ANTHROPIC_BASE_URL=https://.../<gateway>
        export ANTHROPIC_AUTH_TOKEN=<token>
    and pass ``--model <name>``.

Example::

    python claude_code_agent.py \
        --claude /path/to/claude --model IQuest-Coder \
        --task-file ../dataset/dataset.xlsx \
        --source-dir ../dataset/task_sheets \
        --save-path ./claude_code_results --workers 4
"""

import argparse
import concurrent.futures as cf
import datetime
import json
import os
import shutil
import subprocess
import sys

import pandas as pd
import yaml

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "utils"))
import importlib.util


def _load(name, relpath):
    spec = importlib.util.spec_from_file_location(name, os.path.join(HERE, relpath))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod

TrajectoryLogger = _load("sc_traj", "utils/trajectory.py").TrajectoryLogger


PROMPT_TEMPLATE = """You are completing ONE Microsoft Excel task by editing a workbook on Linux (there is no Excel; use Python with openpyxl/pandas, or headless LibreOffice `soffice`).

Files in the current working directory:
- `{source}`  : the original workbook (reference; do not modify).
- `{result}`  : your WORKING copy (already identical to the source). Edit THIS file and save the finished workbook back to this exact path.

Task context: {context}
Task instruction: {instruction}

Requirements (important for automated grading):
- Fulfil the instruction by modifying `{result}` in place and saving it as .xlsx at the same path.
- Keep existing data and sheets intact unless the task says to change them.
- If you add a new sheet, insert it to the LEFT of the first existing sheet.
- Every sheet's table must start at cell A1 and be contiguous (no leading blank rows/cols).
- The grader reads the CACHED cell values, and openpyxl does NOT evaluate formulas. So whenever the task needs a computed result, make sure real values are stored: either compute them in Python and write the values, OR after writing formulas recalculate by converting through LibreOffice, e.g.
    `soffice --headless --calc --convert-to xlsx --outdir <dir> {result}` (then move the recalculated file back to `{result}`).
- Work fully autonomously; do NOT ask questions. When finished, briefly state what you changed.
"""


def build_claude_env(args):
    """Environment for the Claude Code CLI, mirroring the working Harbor setup.

    Two things matter when pointing Claude Code at a custom (non-Anthropic)
    gateway via ``ANTHROPIC_BASE_URL``:

    * route *every* model alias (main + the small/fast "haiku" model Claude Code
      uses for background tasks + subagents) to the same backend model, otherwise
      those auxiliary calls hit model names the gateway doesn't serve;
    * for **DeepSeek-V4 thinking models**, pass
      ``CLAUDE_CODE_EXTRA_BODY={"chat_template_kwargs":{"thinking":true,"reasoning_effort":true}}``.
      Without it the gateway enforces "reasoning_content must be passed back in
      thinking mode" and rejects every multi-turn follow-up (after the first tool
      call), so the agent can never edit the workbook.
    """
    env = dict(os.environ)
    env["CLAUDE_CODE_DISABLE_NONESSENTIAL_TRAFFIC"] = "1"
    model = args.model
    env["ANTHROPIC_MODEL"] = model
    if env.get("ANTHROPIC_BASE_URL"):
        env["ANTHROPIC_DEFAULT_SONNET_MODEL"] = model
        env["ANTHROPIC_DEFAULT_OPUS_MODEL"] = model
        env["ANTHROPIC_DEFAULT_HAIKU_MODEL"] = model
        env["CLAUDE_CODE_SUBAGENT_MODEL"] = model

    extra_body = args.extra_body
    if extra_body is None and "deepseek-v4" in model.lower():
        extra_body = '{"chat_template_kwargs":{"thinking":true,"reasoning_effort":true}}'
    if extra_body:
        env["CLAUDE_CODE_EXTRA_BODY"] = extra_body
    return env


def task_name(index, row):
    return f"{index + 1}_{row['Sheet Name']}"


def build_prompt(row, source_name, result_name):
    return PROMPT_TEMPLATE.format(
        source=source_name,
        result=result_name,
        context=str(row.get("Context", "") or ""),
        instruction=str(row.get("Instructions", "") or ""),
    )


def run_one(args, index, row):
    name = task_name(index, row)
    task_dir = os.path.join(args.save_path, name)
    os.makedirs(task_dir, exist_ok=True)

    src = os.path.join(args.source_dir, f"{row['Sheet Name']}.xlsx")
    source_name = f"{name}_source.xlsx"
    result_name = f"{name}_1.xlsx"
    source_path = os.path.join(task_dir, source_name)
    result_path = os.path.join(task_dir, result_name)
    log_path = os.path.join(task_dir, f"{name}_log.yaml")

    if not os.path.exists(source_path):
        shutil.copy(src, source_path)
    # Start the working result as an exact copy of the source.
    shutil.copy(src, result_path)

    context = str(row.get("Context", "") or "")
    instruction = str(row.get("Instructions", "") or "")
    prompt = build_prompt(row, source_name, result_name)

    cmd = [
        args.claude, "-p", prompt,
        "--model", args.model,
        "--permission-mode", args.permission_mode,
        "--output-format", "json",
        "--max-turns", str(args.max_turns),
    ]
    if args.allowed_tools:
        cmd += ["--allowedTools", args.allowed_tools]

    env = build_claude_env(args)

    t0 = datetime.datetime.now()
    error = None
    result_json = None
    final_text = ""
    num_turns = None
    cost = None
    usage = None
    try:
        proc = subprocess.run(
            cmd, cwd=task_dir, env=env, stdin=subprocess.DEVNULL,
            capture_output=True, text=True, timeout=args.timeout,
        )
        out = proc.stdout.strip()
        try:
            result_json = json.loads(out)
            final_text = result_json.get("result") or result_json.get("text") or ""
            num_turns = result_json.get("num_turns")
            cost = result_json.get("total_cost_usd")
            u = result_json.get("usage") or {}
            if u:
                usage = {
                    "prompt_tokens": u.get("input_tokens"),
                    "completion_tokens": u.get("output_tokens"),
                    "total_tokens": (u.get("input_tokens") or 0) + (u.get("output_tokens") or 0),
                }
        except Exception:
            final_text = out
        if proc.returncode != 0:
            error = (proc.stderr or "")[:2000] or f"claude exited {proc.returncode}"
    except subprocess.TimeoutExpired:
        error = f"claude timed out after {args.timeout}s"
    except Exception as e:
        error = str(e)

    latency = (datetime.datetime.now() - t0).total_seconds()

    # Did the result actually change vs. the source?
    changed = _files_differ(source_path, result_path)
    success = (error is None) and changed and os.path.exists(result_path)

    # ---- evaluator-compatible log.yaml ----
    # refined response drives action counting (A50/A90); one entry per Claude turn.
    refined = [["claude_turn"]] * (num_turns or 1)
    response_record = {
        "refined response": refined,
        "claude_result": final_text[:4000],
        "num_turns": num_turns,
        "cost_usd": cost,
        "changed_workbook": changed,
        "error": error,
    }
    log = {
        "Source Path": os.path.abspath(source_path),
        "Context": context,
        "Instructions": instruction,
        "Success Response": [response_record] if success else [],
        "Fail Response": [] if success else [response_record],
        "Success Count": 1 if success else 0,
        "Total Count": 1,
        "Prompt_format": "claude-code",
        "Agent": "claude-code",
        "Model": args.model,
        "Timestamp": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    with open(log_path, "w") as f:
        yaml.dump(log, f, allow_unicode=True)

    # ---- trajectory.json (for visualize_trajectories.py) ----
    traj = TrajectoryLogger(meta={
        "task": name, "model": args.model, "agent": "claude-code",
        "instruction": instruction, "context": context,
        "source_file": os.path.abspath(source_path),
        "result_file": os.path.abspath(result_path),
        "mode": "claude_code",
    })
    traj.record_call(
        stage="claude_code", model=args.model,
        request_messages=[{"role": "user", "content": prompt}],
        response_content=final_text,
        usage=usage, latency_s=latency, error=error,
        extra={"num_turns": num_turns, "cost_usd": cost, "changed_workbook": changed},
    )
    res_dir = os.path.join(task_dir, result_name.replace(".xlsx", ""))
    traj.save(os.path.join(res_dir, f"{result_name.replace('.xlsx', '')}_trajectory.json"),
              success=success)

    flag = "ok " if success else "!! "
    print(f"  {flag}{name:<28} turns={num_turns} changed={changed} {latency:5.0f}s"
          + (f"  ERROR: {error[:60]}" if error else ""))
    return {"task": name, "success": success, "changed": changed, "num_turns": num_turns,
            "cost_usd": cost, "latency_s": round(latency, 1), "error": error}


def _files_differ(a, b):
    try:
        import openpyxl
        wa = openpyxl.load_workbook(a, data_only=True)
        wb = openpyxl.load_workbook(b, data_only=True)
        if len(wa.worksheets) != len(wb.worksheets):
            return True
        for sa, sb in zip(wa.worksheets, wb.worksheets):
            if sa.max_row != sb.max_row or sa.max_column != sb.max_column:
                return True
            for r in range(1, (sa.max_row or 0) + 1):
                for c in range(1, (sa.max_column or 0) + 1):
                    if sa.cell(r, c).value != sb.cell(r, c).value:
                        return True
        return False
    except Exception:
        # Fall back to a byte comparison.
        try:
            return open(a, "rb").read() != open(b, "rb").read()
        except Exception:
            return True


def main():
    p = argparse.ArgumentParser(description="Run Claude Code as a SheetCopilot agent over the task set.")
    p.add_argument("--claude", default=os.environ.get("CLAUDE_BIN", "claude"), help="path to the claude binary")
    p.add_argument("--model", default=os.environ.get("ANTHROPIC_MODEL", "IQuest-Coder"))
    p.add_argument("--task-file", default=os.path.join(HERE, "..", "dataset", "dataset.xlsx"))
    p.add_argument("--source-dir", default=os.path.join(HERE, "..", "dataset", "task_sheets"))
    p.add_argument("--save-path", default=os.path.join(HERE, "claude_code_results"))
    p.add_argument("--num", type=int, default=None, help="limit to the first N tasks (default: all)")
    p.add_argument("--workers", type=int, default=4)
    p.add_argument("--max-turns", type=int, default=30)
    p.add_argument("--timeout", type=int, default=900, help="per-task timeout (seconds)")
    p.add_argument("--permission-mode", default="default")
    p.add_argument("--allowed-tools", default="Bash Read Write Edit Glob Grep",
                   help="tools to pre-allow so they run without prompting in -p mode")
    p.add_argument("--extra-body", default=None,
                   help="JSON for CLAUDE_CODE_EXTRA_BODY (auto-set for DeepSeek-V4 thinking models)")
    args = p.parse_args()

    os.makedirs(args.save_path, exist_ok=True)
    df = pd.read_excel(args.task_file, header=0)
    if args.num:
        df = df.head(args.num)

    # Checkpoint: skip tasks already done (log.yaml present).
    todo = []
    for index, row in df.iterrows():
        log_path = os.path.join(args.save_path, task_name(index, row), f"{task_name(index, row)}_log.yaml")
        if os.path.exists(log_path):
            continue
        todo.append((index, row))

    print(f"Claude Code agent: model={args.model}  tasks={len(todo)}/{len(df)} (rest cached)")
    print(f"  claude={args.claude}  save_path={args.save_path}  workers={args.workers}\n")

    results = []
    with cf.ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {ex.submit(run_one, args, i, r): i for i, r in todo}
        for fut in cf.as_completed(futs):
            results.append(fut.result())

    ok = sum(1 for r in results if r["success"])
    print(f"\nDone. {ok}/{len(results)} produced a modified workbook.")
    print(f"Now evaluate with:\n  python evaluation_ubuntu.py -c <config with save_path={args.save_path}>")
    summary = {
        "model": args.model, "save_path": os.path.abspath(args.save_path),
        "completed": len(results), "produced_workbook": ok, "results": results,
    }
    with open(os.path.join(args.save_path, "claude_code_run_summary.json"), "w") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    main()
