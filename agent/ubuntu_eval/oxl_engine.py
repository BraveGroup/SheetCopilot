"""
openpyxl-based comparison engine for the Ubuntu/Linux port of the SheetCopilot
outcome-based evaluator.

This module faithfully reproduces the *cell*, *conditional-formatting*, *filter*
and *view (frozen panes)* comparisons implemented for Windows in
``agent/utils/compare_sheets.py`` (which drives Excel through ``win32com``),
but using only ``openpyxl`` so that it runs on Ubuntu without Excel.

Key faithfulness principle
--------------------------
Both the ground-truth workbook and the result workbook are ``.xlsx`` files that
were produced/saved by Excel.  Therefore *any consistent* extraction performed
by openpyxl on both files yields the same relative comparison as the original
Excel-COM code, even when the absolute representation differs (e.g. raw RGB vs.
``ColorIndex`` buckets, or ``int`` vs ``float`` for numbers).  Where the original
code has observable quirks that influence the pass/fail decision (for instance
the cell *data type* is always compared once values match, regardless of whether
``formatting`` was requested) we deliberately reproduce them so the metrics match
the Windows evaluator.

Charts and pivot tables are intentionally *not* handled here -- they require
formula/series recomputation and are routed to the LibreOffice-UNO engine
(``uno_engine.py``) by the dispatcher.
"""

import numpy as np
import openpyxl

# Tolerance used by the original evaluator for numeric comparisons.
TOLERANCE = 1e-8


def _patch_openpyxl_custom_filter():
    """Relax openpyxl's overly strict AutoFilter ``customFilter`` value check.

    A few dataset workbooks store custom-filter criteria (e.g. plain text or
    mid-string patterns) that Excel and LibreOffice accept but that openpyxl
    rejects with "Value must be either numerical or a string containing a
    wildcard", aborting the whole load.  We make the descriptor lenient: store
    any string as a string instead of raising, so the workbook (and its filter
    state) loads intact."""
    try:
        from openpyxl.worksheet.filters import CustomFilterValueDescriptor, Convertible

        def _lenient_set(self, instance, value):
            if isinstance(value, str):
                self.expected_type = str
            Convertible.__set__(self, instance, value)

        CustomFilterValueDescriptor.__set__ = _lenient_set
    except Exception:
        pass


_patch_openpyxl_custom_filter()

# Excel sentinels for ``Range.End`` (the original uses ``A1.End(xlDown)`` and
# ``A1.End(xlToRight)``).  When a contiguous block cannot be found the COM call
# walks to the worksheet boundary; we emulate that with these limits and then
# clamp with ``UsedRange`` exactly like the original ``min(...)`` expressions.
XL_MAX_ROW = 1048576
XL_MAX_COL = 16384


# --------------------------------------------------------------------------- #
# Cached-value workbook loading
# --------------------------------------------------------------------------- #
def load_values_workbook(path):
    """Load a workbook exposing Excel's *cached* formula results.

    ``data_only=True`` makes ``cell.value`` return the value Excel last computed
    and stored in the file (mirroring COM ``Range.Value``).  Styles, number
    formats, conditional formatting, auto-filters and frozen panes are all still
    available on this same object, so a single load serves every openpyxl check.
    """
    return openpyxl.load_workbook(path, data_only=True)


# --------------------------------------------------------------------------- #
# Value/format helper signatures
# --------------------------------------------------------------------------- #
def _norm_number(v):
    """Coerce Excel numbers to ``float``.

    Excel stores every number as an IEEE double, so COM always returned floats;
    openpyxl may return ``int`` when the stored text has no decimal point.  The
    original code mismatches on ``type(a) != type(b)`` -- to keep that semantic
    meaningful we normalise ints to floats first (booleans are left untouched,
    since Excel booleans are a distinct type)."""
    if isinstance(v, bool):
        return v
    if isinstance(v, int):
        return float(v)
    return v


def _color_sig(color):
    """A hashable signature for an openpyxl ``Color`` (font/fill colour)."""
    if color is None:
        return None
    t = getattr(color, "type", None)
    if t == "rgb":
        return ("rgb", color.rgb)
    if t == "theme":
        tint = getattr(color, "tint", 0.0) or 0.0
        return ("theme", color.theme, round(tint, 4))
    if t == "indexed":
        return ("indexed", color.indexed)
    return (t, getattr(color, "value", None))


def _font_sig(font):
    """Signature mirroring the Font attributes the original compares:
    Name, Size, Color, Bold, Italic, Underline."""
    if font is None:
        return None
    return (
        font.name,
        font.size,
        _color_sig(font.color),
        bool(font.bold),
        bool(font.italic),
        font.underline,
    )


def _fill_sig(fill):
    """Signature for the cell interior (``Interior.Color`` in COM)."""
    if fill is None:
        return None
    return (
        getattr(fill, "patternType", None),
        _color_sig(getattr(fill, "fgColor", None)),
        _color_sig(getattr(fill, "bgColor", None)),
    )


def get_datatype(number_format):
    """Port of ``compare_sheets.get_datatype`` -- maps an Excel number-format
    string to one of the coarse data-type buckets the original uses."""
    lower_number_format = (number_format or "").lower()
    if lower_number_format in ["currency"] or "$" in lower_number_format or "¥" in lower_number_format:
        return "currency"
    elif "%" in lower_number_format:
        return "percentage"
    elif "yy" in lower_number_format or "dd" in lower_number_format:
        return "date"
    elif any([x in lower_number_format for x in [":", "mm", "m", "ss", "hh", "am", "pm"]]):
        return "time"
    elif lower_number_format == "@":
        return "text"
    elif "general" in lower_number_format:
        return "general"
    elif "$" not in lower_number_format and "¥" not in lower_number_format and (
        "#0" in lower_number_format or lower_number_format in ["0", "0.00"]
    ):
        return "number"
    else:
        return number_format


# --------------------------------------------------------------------------- #
# Used-range / contiguous-block dimensions (emulating Range.End)
# --------------------------------------------------------------------------- #
def _cell_val(ws, row, col):
    return ws.cell(row=row, column=col).value


def _end_down(ws, col, max_row):
    """Emulate ``Range(A1).End(xlDown).Row`` walking down ``col`` from row 1."""
    if _cell_val(ws, 2, col) is None:
        r = 2
        while r <= max_row and _cell_val(ws, r, col) is None:
            r += 1
        return r if r <= max_row else XL_MAX_ROW
    r = 1
    while r + 1 <= max_row and _cell_val(ws, r + 1, col) is not None:
        r += 1
    return r


def _end_right(ws, row, max_col):
    """Emulate ``Range(A1).End(xlToRight).Column`` walking right along ``row``."""
    if _cell_val(ws, row, 2) is None:
        c = 2
        while c <= max_col and _cell_val(ws, row, c) is None:
            c += 1
        return c if c <= max_col else XL_MAX_COL
    c = 1
    while c + 1 <= max_col and _cell_val(ws, row, c + 1) is not None:
        c += 1
    return c


def get_dims(ws):
    """Return ``(NumberOfRows, NumberOfColumns)`` exactly as the original:
    ``min(A1.End(xlDown).Row, UsedRange.Rows.Count)`` and the column analogue.
    ``UsedRange`` is approximated by ``max_row``/``max_column`` which is accurate
    for the contiguous-from-A1 tables this benchmark mandates."""
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    nrows = min(_end_down(ws, 1, max_row), max_row)
    ncols = min(_end_right(ws, 1, max_col), max_col)
    return nrows, ncols


def _read_block(ws, nrows, ncols):
    """Return a 2-D list (row-major, 1-based emulated by index 0) of cell values
    for the top-left ``nrows`` x ``ncols`` block."""
    block = []
    for r in range(1, nrows + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ncols + 1)]
        block.append(row_vals)
    return block


# --------------------------------------------------------------------------- #
# Cell comparison (port of compare_cells_itercolumn)
# --------------------------------------------------------------------------- #
def compare_cells(ws1, ws2, report):
    """Header-matched, column-by-column cell comparison.

    ``ws1`` is the ground truth, ``ws2`` the result.  Mirrors
    ``compare_sheets.compare_cells_itercolumn`` including: matching result
    columns to GT columns by their header text, the every-100-rows sampling for
    tall tables, the type-aware value comparison, and the quirk that the cell
    *data type* (number-format bucket) is always enforced once values match."""
    cells_report = report.get("cells", {})
    check_values = "values" in cells_report
    check_formatting = "formatting" in cells_report
    check_hyperlink = "hyperlink" in cells_report

    nrows, ncols = get_dims(ws1)
    nrows2, ncols2 = get_dims(ws2)

    # Dimension mismatch -> values fail (unconditionally, as in the original).
    if (nrows, ncols) != (nrows2, ncols2):
        report["cells"]["values"] = 0
        return

    block1 = _read_block(ws1, nrows, ncols)
    block2 = _read_block(ws2, nrows, ncols)

    used_match_col_ids = []
    mismatch = False

    for col in range(ncols):
        header_ws1 = block1[0][col]

        # Find the (still unused) result column whose header equals the GT header.
        match_col_id = 0
        while match_col_id < ncols:
            if header_ws1 == block2[0][match_col_id] and match_col_id not in used_match_col_ids:
                used_match_col_ids.append(match_col_id)
                break
            match_col_id += 1
        else:
            report["cells"]["values"] = 0
            return

        step = 100 if nrows > 200 else 1

        for row in range(0, nrows, step):
            c1 = ws1.cell(row=row + 1, column=col + 1)
            c2 = ws2.cell(row=row + 1, column=match_col_id + 1)
            v1 = block1[row][col]
            v2 = block2[row][match_col_id]

            # ---- value comparison (type-aware) ----
            if check_values:
                n1, n2 = _norm_number(v1), _norm_number(v2)
                cell_mismatch = False
                if type(n1) != type(n2):
                    cell_mismatch = True
                elif isinstance(n1, str) and isinstance(n2, str):
                    if n1.strip() != n2.strip():
                        cell_mismatch = True
                elif isinstance(n1, (int, float)) and isinstance(n2, (int, float)):
                    if not np.allclose(n1, n2, atol=TOLERANCE):
                        cell_mismatch = True
                elif not isinstance(n1, (int, float, str)) and not isinstance(n2, (int, float, str)):
                    if n1 != n2:
                        cell_mismatch = True

                if cell_mismatch:
                    report["cells"]["values"] = 0
                    mismatch = True
                    break

            # ---- formatting + data type ----
            # Reproduce the original precedence `(formatting and <fmt diff>) or <datatype diff>`,
            # i.e. the data-type bucket is always enforced for value-matching cells.
            fmt_diff = False
            if check_formatting:
                fmt_diff = (
                    _font_sig(c1.font) != _font_sig(c2.font)
                    or _fill_sig(c1.fill) != _fill_sig(c2.fill)
                )
            datatype_diff = get_datatype(c1.number_format) != get_datatype(c2.number_format)
            if (check_formatting and fmt_diff) or datatype_diff:
                report["cells"]["formatting"] = 0
                mismatch = True
                break

            # ---- hyperlinks (headers only, rows <= 2 like the original) ----
            if check_hyperlink and row <= 2:
                h1 = c1.hyperlink.target if c1.hyperlink is not None else None
                h2 = c2.hyperlink.target if c2.hyperlink is not None else None
                if h1 != h2:
                    report["cells"]["hyperlink"] = 0
                    mismatch = True
                    break

        if mismatch:
            break


# --------------------------------------------------------------------------- #
# Conditional formatting (port of compare_format_conditions + its caller)
# --------------------------------------------------------------------------- #
def _iter_cf_rules(ws):
    """Yield every conditional-formatting rule on the worksheet."""
    rules = []
    cf = ws.conditional_formatting
    for rng in cf:
        for rule in cf[rng]:
            rules.append(rule)
    return rules


def _cf_formula1(rule):
    f = getattr(rule, "formula", None)
    if f:
        return f[0]
    return None


def _cf_font(rule):
    dxf = getattr(rule, "dxf", None)
    if dxf is None:
        return None
    return getattr(dxf, "font", None)


def _cf_fill(rule):
    dxf = getattr(rule, "dxf", None)
    if dxf is None:
        return None
    return getattr(dxf, "fill", None)


def _compare_one_cf(rule1, rule2, props):
    """Per-pair conditional-format comparison.  ``props`` is the per-property
    report dict for ``format_conditions`` (keys present == requested)."""
    out = dict(props)
    f1, f2 = _cf_font(rule1), _cf_font(rule2)

    if "formula1" in out:
        if rule1.type != rule2.type or _cf_formula1(rule1) != _cf_formula1(rule2):
            out["formula1"] = 0
    if "color" in out:
        c1 = _color_sig(getattr(f1, "color", None)) if f1 is not None else None
        c2 = _color_sig(getattr(f2, "color", None)) if f2 is not None else None
        if c1 != c2:
            out["color"] = 0
    if "fill_color" in out:
        if _fill_sig(_cf_fill(rule1)) != _fill_sig(_cf_fill(rule2)):
            out["fill_color"] = 0
    if "font" in out:
        n1 = getattr(f1, "name", None) if f1 is not None else None
        n2 = getattr(f2, "name", None) if f2 is not None else None
        if n1 != n2:
            out["font"] = 0
    if "bold" in out:
        b1 = bool(getattr(f1, "bold", False)) if f1 is not None else False
        b2 = bool(getattr(f2, "bold", False)) if f2 is not None else False
        if b1 != b2:
            out["bold"] = 0
    if "italic" in out:
        i1 = bool(getattr(f1, "italic", False)) if f1 is not None else False
        i2 = bool(getattr(f2, "italic", False)) if f2 is not None else False
        if i1 != i2:
            out["italic"] = 0
    if "underline" in out:
        u1 = getattr(f1, "underline", None) if f1 is not None else None
        u2 = getattr(f2, "underline", None) if f2 is not None else None
        if u1 != u2:
            out["underline"] = 0
    return out


def compare_format_conditions(ws1, ws2, report):
    """Port of the ``format_conditions`` branch of ``compare_worksheets``."""
    props = report.get("format_conditions", {})
    rules1 = _iter_cf_rules(ws1)
    rules2 = _iter_cf_rules(ws2)

    if len(rules1) != len(rules2):
        report["format_conditions"]["count"] = 0
        return
    if len(rules1) == 0 and len(rules2) == 0:
        for k in report["format_conditions"]:
            report["format_conditions"][k] = 1
        return

    pairs = [(r1, r2) for r1 in rules1 for r2 in rules2]
    last = None
    for r1, r2 in pairs:
        last = _compare_one_cf(r1, r2, props)
        if all(last.values()):
            for k in report["format_conditions"]:
                report["format_conditions"][k] = 1
            return

    # No fully matching pair found.
    if len(pairs) > 1:
        for k in report["format_conditions"]:
            if k == "count":
                continue
            report["format_conditions"][k] = 0
    else:
        for k in report["format_conditions"]:
            if k == "count":
                continue
            report["format_conditions"][k] = last[k]


# --------------------------------------------------------------------------- #
# Filters (port of compare_filters_by_visible_range)
# --------------------------------------------------------------------------- #
def _visible_filter_values(ws, ref):
    """Collect the values of the cells inside ``ref`` whose row is not hidden."""
    from openpyxl.utils.cell import range_boundaries

    min_col, min_row, max_col, max_row = range_boundaries(ref)
    values = []
    for r in range(min_row, max_row + 1):
        rd = ws.row_dimensions.get(r)
        if rd is not None and rd.hidden:
            continue
        for c in range(min_col, max_col + 1):
            values.append(ws.cell(row=r, column=c).value)
    return values


def _dims_ref(ws):
    """``A1``-anchored used range as an ``A1:.."`` string, for filter fallback."""
    from openpyxl.utils import get_column_letter

    nrows, ncols = get_dims(ws)
    return f"A1:{get_column_letter(ncols)}{nrows}"


def compare_filters(ws1, ws2, report):
    """Compare the post-filter *visible* range of GT vs. result.

    The filter outcome is the set of rows left *visible* after AutoFilter, i.e.
    the non-hidden rows of the filtered range.  When a workbook records no
    ``<autoFilter>`` element (a few GTs apply then clear the filter, leaving the
    effect only as hidden/unhidden rows), we fall back to the A1 used range so an
    identical workbook still compares equal."""
    props = report.get("filters", {})
    ref1 = ws1.auto_filter.ref or _dims_ref(ws1)
    ref2 = ws2.auto_filter.ref or _dims_ref(ws2)

    vis1 = _visible_filter_values(ws1, ref1)
    vis2 = _visible_filter_values(ws2, ref2)
    if len(vis1) != len(vis2) or any(a != b for a, b in zip(vis1, vis2)):
        for k in props:
            report["filters"][k] = 0


# --------------------------------------------------------------------------- #
# View / frozen panes (port of compare_frozen_panes)
# --------------------------------------------------------------------------- #
def _freeze_split(ws):
    """Return ``(frozen_rows, frozen_cols)`` from ``ws.freeze_panes``.

    ``freeze_panes == 'B2'`` means 1 row + 1 column are frozen, matching the COM
    ``ActiveWindow.SplitRow`` / ``SplitColumn`` the original reads."""
    fp = ws.freeze_panes
    if not fp:
        return (0, 0)
    from openpyxl.utils.cell import coordinate_to_tuple

    row, col = coordinate_to_tuple(fp)
    return (row - 1, col - 1)


def compare_view(ws1, ws2, report):
    """Compare frozen-pane configuration."""
    if "freeze_pane" not in report.get("view", {}):
        return
    if _freeze_split(ws1) != _freeze_split(ws2):
        report["view"]["freeze_pane"] = 0
